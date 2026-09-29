from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
import os
import pandas as pd

# Perdida de Carga
def csv2xlsx_PerdidaDeCarga(outputFile:str,title:str,df:pd.DataFrame,df_temporal:pd.DataFrame):
    print("Entra en csv2xlsx Perdida de Carga")
    def font_editor(tabla, fila, columna, dato, font_size = 14, negrita = False, centrado_vertical = True, centrado_horizontal = True):
            cell = tabla.cell(fila,columna)

            p = cell.paragraphs[0]
            p.text = ""
            run = p.add_run(str(dato))

            run.bold = negrita
            run.font.size = Pt(font_size)

            if centrado_vertical:
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            
            if centrado_horizontal:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def add_run_header(p, texto, negrita = True, subrallado = True, font_size = 18):
            run = p.add_run(texto)
            run.bold = negrita
            run.underline = subrallado
            run.font.size = Pt(font_size)

    template_word = os.path.join("GuardadoDatos","Templates","template_perdida_de_carga.docx")
    # Guardado a Word
    doc = Document(template_word)
    section = doc.sections[0]
    header = section.header
    
    for p in header.paragraphs:
        for run in p.runs:
            p.run = ""
        
        add_run_header(p,"Pérdida de Carga:")
        add_run_header(p," "*10, subrallado=False)
        add_run_header(p,"FECHA:")
        add_run_header(p," " + title[0:2]+"/"+title[2:4]+"/"+title[4:6]+ " "*8, subrallado=False)
        add_run_header(p,"N° de Ensayo:")
        add_run_header(p," "+title+" "*6, subrallado= False)
        add_run_header(p,"Hoja N°:")
        add_run_header(p," "*2 + "de" + " "*3,subrallado= False)
        break
    
    ubicacion_datos_word ={
            "Caudal": 0,
            "Ps": 1,
            "Pe": 2,
            "Ps-Pe": 3,
        }

    talba_datos = doc.tables[0]

    font_editor(talba_datos,fila = 1, columna = 4, dato =str(df_temporal["T_amb"].iloc[0]), negrita = True)
    font_editor(talba_datos,fila = 1, columna = 8, dato =str(df_temporal["T_amb"].iloc[-1]), negrita = True)
    font_editor(talba_datos,fila = 2, columna = 4, dato =str(df_temporal["T_H2O"].iloc[0]), negrita = True)
    font_editor(talba_datos,fila = 2, columna = 8, dato =str(df_temporal["T_H2O"].iloc[-1]), negrita = True)
    font_editor(talba_datos,fila = 3, columna = 4, dato =str(df_temporal["Horario"].iloc[0]), negrita = True)
    font_editor(talba_datos,fila = 3, columna = 8, dato =str(df_temporal["Horario"].iloc[-1]), negrita = True)
    i_fila = 0
    i_col = 0
    n_rows = 8
    fila_start = 5
    for key, value in ubicacion_datos_word.items():
        if key in df.columns:
            valores = df[key].tolist()
            for i, val in enumerate(valores):
                i_fila = i -n_rows*(i//n_rows) + fila_start
                # print(f"i_fila: {i_fila}")
                i_col = 4*(i//n_rows)
                # print(f"i_col: {i_col}")
                font_editor(talba_datos,fila= i_fila,columna = i_col + value, dato = val)
    
    output_file_word = outputFile[:-5] + ".docx"
    print(f"Guardado en formato word como {output_file_word}, COMPLETAR")
    doc.save(output_file_word)